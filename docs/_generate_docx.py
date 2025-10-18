from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

# Create docx
doc = Document()
# Set default font (Normal style) to Times New Roman, 12pt
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
# Ensure East Asian font setting for some Word viewers
rfonts = style.element.rPr.rFonts
rfonts.set(qn('w:eastAsia'), 'Times New Roman')
doc.add_heading('Testing', level=1)

doc.add_paragraph("This section outlines testing performed for the VendWise mobile application. It includes unit, integration, and manual test cases. The table below is formatted for inclusion in reports and papers; fill in the Actual Output and Pass/Fail columns during execution.")

doc.add_heading('Testing strategy', level=2)
doc.add_paragraph('- Unit tests: small logical units and helpers.')
doc.add_paragraph('- Widget tests: verify UI widgets behavior in isolation.')
doc.add_paragraph('- Integration tests: end-to-end flows involving Supabase (auth, transactions, sync).')
doc.add_paragraph('- Manual tests: sensor permissions, camera, offline sync, and UX checks.')

# Table rows
rows = [
    ['Test Case ID','Test Case Description','Input','Expected Output','Actual Output','Pass/Fail','Notes'],
    ['TC-01','User sign-in with valid credentials','Email/password','Successful login and navigation to dashboard','Not executed - requires Supabase credentials','Pending','Manual test / integration environment required'],
    ['TC-02','User sign-in with invalid credentials','Wrong password','Error message shown, no login','Not executed - requires Supabase credentials','Pending','Manual test / integration environment required'],
    ['TC-03','Create inventory item','Name, SKU, price','Item appears in inventory list','Not executed - requires interactive UI test','Pending','Manual test on device or emulator required'],
    ['TC-04','Create sale transaction (online)','Product, qty, price','Transaction recorded in DB and shown in dashboard top-10','Not executed - requires Supabase and UI','Pending','Manual/integration test required'],
    ['TC-05','Create transaction while offline','Product, qty, price (offline)','Transaction saved locally and synced when online','Not executed - offline test not run','Pending','Manual test: toggle network on device/emulator'],
    ['TC-06','Recent transactions inline limit','Many transactions (>10)','Dashboard/Report shows only 10 items inline; "View More" loads full list','Not executed - UI verification needed','Pending','Manual test on dashboard/reports screens'],
    ['TC-07','Dismiss a notification','Notification displayed; user taps dismiss','Notification no longer appears in future sessions','Not executed - interactive UI test','Pending','Manual test: dismiss and restart app to verify persistence'],
    ['TC-08','Camera permission & barcode scan','Camera permission grant + barcode image','Product scanned and filled in form','Not executed - device camera required','Pending','Manual test on a device with camera'],
    ['TC-09','Build Android release','Build command','AAB produced at build/app/outputs/bundle/release/app-release.aab','AAB exists at build/app/outputs/bundle/release/app-release.aab','Pass','Built earlier in session; check file path on disk'],
    ['TC-10','iOS archive and export','Build/archive in Xcode','Archive created; IPA export requires Apple signing','Archive exists at build/ios/archive/Runner.xcarchive; IPA export requires Apple signing','Partial (Archive created)','Requires Apple Distribution certificate / provisioning profile for IPA export'],
]

table = doc.add_table(rows=1, cols=7)
hdr_cells = table.rows[0].cells
for i,heading in enumerate(rows[0]):
    hdr_cells[i].text = heading

for row in rows[1:]:
    cells = table.add_row().cells
    for i,cell in enumerate(row):
        cells[i].text = cell

# Notes
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

def set_normal_font(doc):
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    try:
        rfonts = style.element.rPr.rFonts
        rfonts.set(qn('w:eastAsia'), 'Times New Roman')
    except Exception:
        pass


doc = Document()
set_normal_font(doc)

# --- Testing Section ---
doc.add_heading('Testing', level=1)

doc.add_paragraph("This section outlines testing performed for the VendWise mobile application. It includes unit, integration, and manual test cases. The table below is formatted for inclusion in reports and papers; fill in the Actual Output and Pass/Fail columns during execution.")

doc.add_heading('Testing strategy', level=2)
doc.add_paragraph('- Unit tests: small logical units and helpers.')
doc.add_paragraph('- Widget tests: verify UI widgets behavior in isolation.')
doc.add_paragraph('- Integration tests: end-to-end flows involving Supabase (auth, transactions, sync).')
doc.add_paragraph('- Manual tests: sensor permissions, camera, offline sync, and UX checks.')

# Table rows
rows = [
    ['Test Case ID','Test Case Description','Input','Expected Output','Actual Output','Pass/Fail','Notes'],
    ['TC-01','User sign-in with valid credentials','Email/password','Successful login and navigation to dashboard','Not executed - requires Supabase credentials','Pending','Manual test / integration environment required'],
    ['TC-02','User sign-in with invalid credentials','Wrong password','Error message shown, no login','Not executed - requires Supabase credentials','Pending','Manual test / integration environment required'],
    ['TC-03','Create inventory item','Name, SKU, price','Item appears in inventory list','Not executed - requires interactive UI test','Pending','Manual test on device or emulator required'],
    ['TC-04','Create sale transaction (online)','Product, qty, price','Transaction recorded in DB and shown in dashboard top-10','Not executed - requires Supabase and UI','Pending','Manual/integration test required'],
    ['TC-05','Create transaction while offline','Product, qty, price (offline)','Transaction saved locally and synced when online','Not executed - offline test not run','Pending','Manual test: toggle network on device/emulator'],
    ['TC-06','Recent transactions inline limit','Many transactions (>10)','Dashboard/Report shows only 10 items inline; "View More" loads full list','Not executed - UI verification needed','Pending','Manual test on dashboard/reports screens'],
    ['TC-07','Dismiss a notification','Notification displayed; user taps dismiss','Notification no longer appears in future sessions','Not executed - interactive UI test','Pending','Manual test: dismiss and restart app to verify persistence'],
    ['TC-08','Camera permission & barcode scan','Camera permission grant + barcode image','Product scanned and filled in form','Not executed - device camera required','Pending','Manual test on a device with camera'],
    ['TC-09','Build Android release','Build command','AAB produced at build/app/outputs/bundle/release/app-release.aab','AAB exists at build/app/outputs/bundle/release/app-release.aab','Pass','Built earlier in session; check file path on disk'],
    ['TC-10','iOS archive and export','Build/archive in Xcode','Archive created; IPA export requires Apple signing','Archive exists at build/ios/archive/Runner.xcarchive; IPA export requires Apple signing','Partial (Archive created)','Requires Apple Distribution certificate / provisioning profile for IPA export'],
]

table = doc.add_table(rows=1, cols=7)
hdr_cells = table.rows[0].cells
for i,heading in enumerate(rows[0]):
    hdr_cells[i].text = heading

for row in rows[1:]:
    cells = table.add_row().cells
    for i,cell in enumerate(row):
        cells[i].text = cell

# Notes
doc.add_paragraph('\nNotes on running tests')
doc.add_paragraph('Run unit/widget tests with:')
doc.add_paragraph('flutter test')

doc.add_paragraph('Run analyzer before committing changes:')
doc.add_paragraph('flutter analyze')

doc.add_paragraph('For integration tests that require Supabase, ensure environment variables (SUPABASE_URL, SUPABASE_KEY) are set in a secure way.')

# --- Implementation Section ---
doc.add_page_break()
set_normal_font(doc)

doc.add_heading('Implementation', level=1)

doc.add_paragraph('Overview')
doc.add_paragraph('The VendWise mobile application is a Flutter client that communicates with a cloud-hosted Supabase backend (Postgres, Auth, Storage). The system adopts an offline-first pattern: the mobile app persists local changes, displays responsive UI immediately, and synchronizes with the backend when network access is available. Authentication, storage, and database are provided by Supabase; application-specific logic lives on the mobile client and can be augmented by server-side functions (Edge Functions) if needed.')

# PlantUML notes (to paste into report or render separately)
doc.add_paragraph('\nDeployment diagram (PlantUML):')
doc.add_paragraph("@startuml\nnode \"Users' Devices\\n(Flutter App)\" as Mobile #LightBlue\ncloud \"Supabase\\n(Auth, Postgres, Storage, Functions)\" as Supabase #LightYellow\nrectangle \"CDN / Static Assets\" as CDN #Lavender\ndatabase \"Postgres (Supabase)\\nPrimary DB\" as DB #LightGray\nqueue \"Background Sync Queue\\n(optional local queue)\" as Queue #LightSteelBlue\n\nMobile --> Supabase : HTTPS (REST / Realtime)\\nAuth tokens\\nSync requests\nSupabase --> DB : SQL read/write\nSupabase --> CDN : Serve/put assets (optional)\nMobile --> CDN : Download media (cached)\nMobile -> Queue : Local operation queue (offline)\nQueue --> Supabase : Push queued operations when online\n\nnote right of Mobile\nOffline-first: local queue + optimistic UI.\nSync uses timestamps / operation log.\nend note\n@enduml")

# Transaction sync sequence
set_normal_font(doc)

doc.add_paragraph('\nTransaction sync sequence (PlantUML):')
doc.add_paragraph("@startuml\nactor User\nparticipant MobileApp\nparticipant LocalDB\nparticipant Supabase\nparticipant Postgres\n\nUser -> MobileApp : Create transaction (UI)\nMobileApp -> LocalDB : Insert transaction (status=queued)\nMobileApp -> MobileApp : Update UI (optimistic)\nMobileApp -> Supabase : (when online) POST /transactions\nSupabase -> Postgres : INSERT transaction\nSupabase --> MobileApp : 201 Created + server id + timestamp\nMobileApp -> LocalDB : Mark transaction as synced (update id, status)\nMobileApp -> User : Show sync success\n@enduml")

# --- Deployment Section ---
doc.add_page_break()
set_normal_font(doc)

doc.add_heading('Deployment', level=1)

doc.add_paragraph('Final rollout: cloud/server installation, user onboarding, release notes.')

doc.add_paragraph('Deployment architecture: Mobile app (Flutter) on Android/iOS devices; Supabase managed backend (Auth, Postgres, Storage); optional CDN for media; CI pipeline for builds and releases. The mobile app follows an offline-first model: it persists local changes and synchronizes with Supabase when connectivity is available.')

doc.add_paragraph('\nConfiguration & environment variables:')
doc.add_paragraph('Do NOT commit secrets to source control. Use CI secret stores (GitHub Actions secrets, GitLab CI variables, or cloud secret managers).')

doc.add_paragraph('\nExample .env (mobile dev / CI):')
doc.add_paragraph('SUPABASE_URL=https://xyzcompany.supabase.co')
doc.add_paragraph('SUPABASE_ANON_KEY=public-anon-key')
doc.add_paragraph('SUPABASE_SERVICE_KEY=service-role-key      # only for server-side (never embed in client)')
doc.add_paragraph('APP_ENV=production')
doc.add_paragraph('SENTRY_DSN=your_sentry_dsn                 # optional for error monitoring')


doc.add_paragraph('\nDeployment checklist:')
doc.add_paragraph('- Prepare Supabase project (DB, Auth, Storage) and enable Row Level Security (RLS) policies')
doc.add_paragraph('- Configure CI secrets for Android keystore and iOS signing certificates')
doc.add_paragraph('- Configure automated backups and monitoring (Supabase backups, Sentry)')
doc.add_paragraph('- Create CI pipeline: analyze -> test -> build -> archive -> upload')
doc.add_paragraph('- Publish Android AAB to Play Console; publish iOS IPA via App Store Connect (requires Apple distribution certs)')

out = 'docs/testing_document_full.docx'
doc.save(out)
print('WROTE', out)
